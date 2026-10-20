"""Excel exports of sellers and items for the admin."""
from io import BytesIO
from zoneinfo import ZoneInfo

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Timestamps are stored in UTC; the sheet shows them the way the desk saw them.
LOCAL_TZ = ZoneInfo("Europe/Berlin")

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EURO_FORMAT = '#,##0.00 "€"'
DATETIME_FORMAT = "DD.MM.YYYY HH:MM"

PAYMENT_LABELS = {"cash": "Bar", "card": "Karte"}


def local(dt):
    """Excel cannot hold timezone-aware datetimes, so convert and drop the zone."""
    if dt is None:
        return None
    return timezone.localtime(dt, LOCAL_TZ).replace(tzinfo=None)


def yes_no(value):
    return "Ja" if value else "Nein"


def item_status(item):
    """Same precedence as the status badge in the item list."""
    if item.is_sold:
        return "Verkauft"
    if item.picked_up_at:
        return "Abgeholt"
    if item.is_stolen:
        return "Gestohlen"
    if not item.accepted_at:
        return "Nicht angenommen"
    if item.returned_at:
        return "Rückgemeldet"
    return "Verfügbar"


def build_workbook(title, columns, rows):
    """columns: list of (header, width, number_format or None)."""
    wb = Workbook()
    ws = wb.active
    ws.title = title

    ws.append([header for header, _, _ in columns])
    header_fill = PatternFill("solid", start_color="4EBDC4")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    for row in rows:
        ws.append(row)

    for index, (_, width, number_format) in enumerate(columns, start=1):
        letter = get_column_letter(index)
        ws.column_dimensions[letter].width = width
        if number_format:
            for cell in ws[letter][1:]:
                cell.number_format = number_format

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return wb


def xlsx_response(wb, name):
    buffer = BytesIO()
    wb.save(buffer)
    filename = f"{name}_{timezone.localtime(timezone.now(), LOCAL_TZ):%Y-%m-%d}.xlsx"
    response = HttpResponse(buffer.getvalue(), content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


SELLER_COLUMNS = [
    ("Nr.", 6, None),
    ("Vorname", 16, None),
    ("Nachname", 18, None),
    ("Straße", 22, None),
    ("Hausnr.", 8, None),
    ("PLZ", 8, None),
    ("Ort", 16, None),
    ("Telefon", 16, None),
    ("Aktive", 8, None),
    ("Großverkäufer", 13, None),
    ("Artikel erfasst", 14, None),
    ("Artikel angenommen", 18, None),
    ("Artikel verkauft", 15, None),
    ("Artikel gestohlen", 16, None),
    ("Umsatz", 12, EURO_FORMAT),
    ("Provision", 12, EURO_FORMAT),
    ("Annahmegebühr", 14, EURO_FORMAT),
    ("Gebühr bezahlt", 14, None),
    ("Auszahlung", 12, EURO_FORMAT),
    ("Angelegt am", 17, DATETIME_FORMAT),
]


def sellers_workbook(sellers):
    rows = []
    for seller in sellers:
        payout = seller.calculate_payout()
        rows.append([
            seller.seller_number,
            seller.first_name,
            seller.last_name,
            seller.street,
            seller.street_number,
            seller.postal_code,
            seller.city,
            seller.mobile_number,
            yes_no(seller.is_member),
            yes_no(seller.is_major_seller),
            len(seller.items.all()),
            payout["total_items_count"],
            payout["sold_items_count"],
            payout["stolen_items_count"],
            payout["total_sales"],
            payout["commission"],
            payout["acceptance_fee"],
            yes_no(seller.acceptance_fee_paid),
            payout["final_payout"],
            local(seller.created_at),
        ])
    return build_workbook("Verkäufer", SELLER_COLUMNS, rows)


ITEM_COLUMNS = [
    ("Artikel-Nr.", 12, None),
    ("Verkäufer-Nr.", 13, None),
    ("Verkäufer", 22, None),
    ("Kategorie", 16, None),
    ("Marke", 16, None),
    ("Farbe", 12, None),
    ("Größe", 10, None),
    ("Preis", 11, EURO_FORMAT),
    ("Beschreibung", 30, None),
    ("Status", 17, None),
    ("Zahlungsart", 12, None),
    ("Angenommen am", 17, DATETIME_FORMAT),
    ("Verkauft am", 17, DATETIME_FORMAT),
    ("Rückgemeldet am", 17, DATETIME_FORMAT),
    ("Abgeholt am", 17, DATETIME_FORMAT),
    ("Angelegt am", 17, DATETIME_FORMAT),
]


def items_workbook(items):
    rows = []
    for item in items:
        # Iterating the prefetched sales keeps this to one query for all items.
        sale = next(iter(item.sales.all()), None)
        rows.append([
            item.barcode,
            item.seller.seller_number,
            f"{item.seller.first_name} {item.seller.last_name}",
            item.category,
            item.brand,
            item.color,
            item.size,
            float(item.price),
            item.description,
            item_status(item),
            PAYMENT_LABELS.get(sale.payment_method) if sale else None,
            local(item.accepted_at),
            local(item.sold_at),
            local(item.returned_at),
            local(item.picked_up_at),
            local(item.created_at),
        ])
    return build_workbook("Artikel", ITEM_COLUMNS, rows)
