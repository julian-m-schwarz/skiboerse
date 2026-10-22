"""Import a major seller's item list from the Excel template the club hands out.

Template columns: Artikelnummer, Verkäufernummer, Artikelbezeichnung, Marke,
Farbe, Größe, Wunschpreis, Kommentar. Columns are found by their header, so
an extra or reordered column does not break the import. The Verkäufernummer
column is ignored: the seller is the one the list is uploaded for.

The seller's own article numbers become the barcode numbers (number 12 of
seller 7 is S007-012), so the labels match the list the seller brought.
"""
import difflib
import re
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from .models import Item

MAX_ARTICLE_NUMBER = 999  # barcodes are S###-###
HEADER_SEARCH_ROWS = 10

# Header text (lowercase, without spaces) → field. Checked with startswith,
# so a header with a remark in brackets still matches.
HEADER_PREFIXES = [
    ("artikelnummer", "number"),
    ("artikelnr", "number"),
    ("artikelbezeichnung", "label"),
    ("bezeichnung", "label"),
    ("kategorie", "label"),
    ("marke", "brand"),
    ("farbe", "color"),
    ("größe", "size"),
    ("grösse", "size"),
    ("groesse", "size"),
    ("wunschpreis", "price"),
    ("preis", "price"),
    ("kommentar", "comment"),
    ("bemerkung", "comment"),
    ("beschreibung", "comment"),
]
REQUIRED_FIELDS = {"label": "Artikelbezeichnung", "price": "Wunschpreis"}

# What sellers write in "Artikelbezeichnung" → item category.
CATEGORY_TERMS = {
    "Snowboard": ["snowboard", "snowboards", "board"],
    "Snowboardboots": [
        "snowboardboots", "snowboardboot", "softboots", "softboot", "boots",
        "snowboardschuhe", "snowboardschuh",
    ],
    "Snowboardbindung": ["snowboardbindung", "snowboardbindungen", "boardbindung"],
    "Ski": ["ski", "skier", "carver", "tourenski"],
    "Skischuhe": ["skischuhe", "skischuh", "skistiefel", "skiboots", "skiboot", "hardboots"],
    "Skibindung": ["skibindung", "skibindungen", "bindung"],
    "Skistoecke": ["skistöcke", "skistoecke", "stöcke", "stoecke", "skistock", "stock"],
    "Helm": ["helm", "skihelm", "snowboardhelm"],
    "Skibrille": ["skibrille", "brille", "goggle", "goggles"],
    "Bekleidung": [
        "bekleidung", "jacke", "skijacke", "hose", "skihose", "anzug", "skianzug",
        "handschuhe", "pullover",
    ],
    "Zubehoer": ["zubehör", "zubehoer", "rucksack", "tasche", "skitasche", "boardbag"],
    "Sonstiges": ["sonstiges"],
}
TERM_TO_CATEGORY = {term: cat for cat, terms in CATEGORY_TERMS.items() for term in terms}
FUZZY_CUTOFF = 0.8


class ImportFileError(Exception):
    """The file as a whole cannot be read; nothing row-specific to report."""


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def match_category(label):
    """Return (category, exact). exact is False for a guess worth showing."""
    normalized = label.lower()
    # "Ski Schuhe" and "Snowboard-Boots" mean the compound word.
    compact = normalized.replace(" ", "").replace("-", "")
    for candidate in (normalized, compact):
        if candidate in TERM_TO_CATEGORY:
            return TERM_TO_CATEGORY[candidate], True

    # "Snowboard mit Bindung" → first word that is a known term.
    for word in normalized.split():
        if word in TERM_TO_CATEGORY:
            return TERM_TO_CATEGORY[word], True

    # Typos such as "Snoboard" or "Softbbots".
    for candidate in [normalized] + normalized.split():
        close = difflib.get_close_matches(candidate, TERM_TO_CATEGORY, n=1, cutoff=FUZZY_CUTOFF)
        if close:
            return TERM_TO_CATEGORY[close[0]], False

    return "Sonstiges", False


def parse_price(value):
    if isinstance(value, (int, float)):
        price = Decimal(str(value))
    else:
        cleaned = text(value).replace("€", "").replace("EUR", "").strip()
        # German notation: 1.234,50
        cleaned = cleaned.replace(".", "").replace(",", ".") if "," in cleaned else cleaned
        price = Decimal(cleaned)
    return price.quantize(Decimal("0.01"))


def parse_size(value):
    if isinstance(value, float) and not value.is_integer():
        return f"{value:g}".replace(".", ",")
    return text(value)


def find_header(rows):
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        columns = {}
        for col, cell in enumerate(row):
            key = text(cell).lower().replace(" ", "")
            for prefix, field in HEADER_PREFIXES:
                if key.startswith(prefix) and field not in columns:
                    columns[field] = col
                    break
        if all(field in columns for field in REQUIRED_FIELDS):
            return index, columns
    raise ImportFileError(
        "Kopfzeile nicht gefunden. Erwartet werden u. a. die Spalten "
        "„Artikelbezeichnung“ und „Wunschpreis“."
    )


def article_number(value):
    """Existing barcode S007-012 → 12."""
    try:
        return int(value.split("-")[1])
    except (IndexError, ValueError):
        return None


def parse_items(file, seller):
    """Read the file and return (items, warnings, errors).

    items are unsaved Item instances with their barcode already assigned.
    Nothing may be saved when errors is not empty.
    """
    try:
        workbook = load_workbook(file, read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("Die Datei ist keine gültige Excel-Datei (.xlsx).")

    sheet = workbook.worksheets[0]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()

    header_index, columns = find_header(rows)

    def cell(row, field):
        col = columns.get(field)
        return row[col] if col is not None and col < len(row) else None

    taken = {
        n for n in (
            article_number(b) for b in Item.objects.filter(seller=seller).values_list("barcode", flat=True)
        ) if n
    }
    content_fields = ["label", "brand", "color", "size", "price", "comment"]

    items, warnings, errors = [], [], []
    numbers_in_file = {}
    pending_number = []  # items without an article number, numbered at the end

    for row_index, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        # Templates often come with numbered but otherwise empty rows.
        if not any(text(cell(row, f)) for f in content_fields):
            continue

        prefix = f"Zeile {row_index}"

        label = text(cell(row, "label"))
        if not label:
            errors.append(f"{prefix}: Artikelbezeichnung fehlt.")
            continue
        category, exact = match_category(label)
        if category == "Sonstiges" and label.lower() != "sonstiges":
            warnings.append(f"{prefix}: „{label}“ nicht erkannt, wird als Sonstiges angelegt.")
        elif not exact:
            warnings.append(f"{prefix}: „{label}“ als {category} erkannt.")

        raw_price = cell(row, "price")
        try:
            price = parse_price(raw_price)
            if price <= 0:
                raise InvalidOperation
        except (InvalidOperation, ValueError):
            errors.append(
                f"{prefix}: Wunschpreis „{text(raw_price)}“ ist ungültig."
                if text(raw_price) else f"{prefix}: Wunschpreis fehlt."
            )
            continue

        item = Item(
            seller=seller,
            name=label,
            category=category,
            brand=text(cell(row, "brand")),
            color=text(cell(row, "color")),
            size=parse_size(cell(row, "size")),
            price=price,
            description=text(cell(row, "comment")),
        )
        for field, limit, title in (("brand", 100, "Marke"), ("color", 50, "Farbe"), ("size", 50, "Größe")):
            if len(getattr(item, field)) > limit:
                errors.append(f"{prefix}: {title} ist länger als {limit} Zeichen.")

        raw_number = text(cell(row, "number"))
        if not raw_number:
            pending_number.append(item)
        elif not raw_number.isdigit() or not 1 <= int(raw_number) <= MAX_ARTICLE_NUMBER:
            errors.append(f"{prefix}: Artikelnummer „{raw_number}“ muss eine Zahl von 1 bis {MAX_ARTICLE_NUMBER} sein.")
            continue
        else:
            number = int(raw_number)
            if number in numbers_in_file:
                errors.append(f"{prefix}: Artikelnummer {number} kommt schon in Zeile {numbers_in_file[number]} vor.")
            elif number in taken:
                errors.append(f"{prefix}: Artikelnummer {number} ist bei diesem Verkäufer schon vergeben.")
            numbers_in_file[number] = row_index
            item.barcode = barcode(seller, number)

        items.append(item)

    # Rows without a number continue after the highest number in use.
    next_number = max(taken | set(numbers_in_file), default=0) + 1
    for item in pending_number:
        if next_number > MAX_ARTICLE_NUMBER:
            errors.append(f"Mehr als {MAX_ARTICLE_NUMBER} Artikel für diesen Verkäufer.")
            break
        item.barcode = barcode(seller, next_number)
        next_number += 1

    if not items and not errors:
        errors.append("Die Datei enthält keine Artikel.")

    return items, warnings, errors


def barcode(seller, number):
    return f"S{seller.seller_number:03d}-{number:03d}"
